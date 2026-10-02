"""levy-check: the user's own checks on a levy roll, for the agent-grid manager.

Reads .agent-grid/checks.toml in the current folder (the folder root). Commands:

  levy-check init               write a starter checks.toml to fill in
  levy-check rules              check the roll the workflow wrote against the rules
  levy-check selftest           plant errors in copies of the roll and last year's certified
                                roll, and report which checks catch each one (changes nothing)
  levy-check spot [FILE]        compare the roll with the user's hand-calculated parcels
                                (FILE, or $AGENT_GRID_SPOT; spawn --spot keeps these outside
                                the folder, where analysts can't read them)
  levy-check backtest-inputs    (used by agent-hook) print "dest<TAB>src" for each input
                                swapped in for the backtest, after a line naming the workflow
  levy-check backtest-compare   compare the roll to last year's certified roll
  levy-check roll-path          (used by agent-hook) print the roll's path

Exit codes: 0 all checks passed, 1 a check failed, 2 checks.toml or a file is unusable.
Amounts are compared exactly, as decimals. APNs are compared without dashes, spaces, or dots.
"""
import csv
import re
import sys
import os
from decimal import ROUND_DOWN, ROUND_HALF_EVEN, ROUND_HALF_UP, ROUND_UP, Decimal, InvalidOperation
from pathlib import Path

try:
    import tomllib
except ImportError:  # Python < 3.11
    tomllib = None

CONFIG = Path(".agent-grid/checks.toml")
EXAMPLES = 10

TEMPLATE = """# Your checks for this folder. The analyst never edits this file (agent-grid blocks it).
# The manager runs them on every test: first on this year's run, then on a backtest that
# reruns the code on last year's inputs and compares to last year's certified roll.
# Paths are relative to the folder root. Files can be .csv or .xlsx.
# Every number here should come from the RMA, the budget, the county data, or last year's
# certified roll, never from the levy code. Note the source next to each one.

[roll]                          # the roll the workflow writes
file = "output/levy_roll.csv"
sheet = ""                      # .xlsx only; blank means the first sheet
apn = "APN"                     # column names, as in the header row
amount = "Levy"
max_tax = ""                    # optional columns, for the rules below
units = ""
rate = ""

[rules]
unique_apn = true               # each APN appears once
unique_by = []                  # or: columns that together identify one line, when a parcel
                                # can appear more than once, e.g. ["APN", "Account"]
no_blank_apn = true
no_negative = true
whole_cents = true              # no fractions of a cent
max_tax = false                 # no amount above the max_tax column
max_tax_is_units_x_rate = false # max_tax equals units x rate, to the cent
rounding = "half_up"            # how cents round: "half_up" (0.125 -> 0.13) or "half_even"
total = 0                       # this year's levy requirement; 0 means don't check
total_tolerance = 1.00
parcels_file = ""               # e.g. the county export: every parcel in it appears in
parcels_apn = "APN"             # the roll, and the roll has no parcel that isn't in it

[totals]                        # optional: expected totals by group (district, zone, fund)
column = ""                     # roll column to group by; blank = no group totals
tolerance = 1.00
[totals.expected]
# "Zone A" = 48250.00           # source: FY 26-27 budget, sheet "Summary", cell D14

[max_rates]                     # optional: the RMA's maximum rate per unit, by rate class.
fiscal_year = ""                # the roll's fiscal year, e.g. "2026-27"
class_column = ""               # roll column naming each parcel's rate class (zone, land use)
units_column = ""               # roll column with each parcel's units; blank = [roll] units,
                                # or 1 per parcel if that's blank too
# Every parcel's levy must be at most its units x its class's max rate. levy-check works
# out the max rate itself, so a wrong escalator in the code can't hide in the roll's own
# columns. The backtest checks last year's roll against last year's max rates too.
# Give each class one of:
# [max_rates.classes."Zone A"]
# rate = 412.3456               # this year's max rate, as a district document states it
# [max_rates.classes."Zone B"]
# rates = { "2025-26" = 300.00, "2026-27" = 306.00 }   # max rate by fiscal year
# [max_rates.classes."Single Family"]
# base_rate = 450.00            # the RMA's base rate ...
# base_year = "2019-20"         # ... in this fiscal year
# escalator = 0.02              # 2% a year, compounded
# places = 4                    # decimals the RMA rounds the max rate to (blank: no rounding)
# rounding = "half_up"          # "half_up", "half_even", or "down"
# round_each_year = false       # true if the RMA rounds every year before escalating again

[backtest]
expected = ""                   # last year's certified roll (the one sent to the county), e.g.
                                # ".agent-grid/backtest/certified_roll.csv"; blank = no backtest
workflow = ""                   # blank = the folder's workflow (if it has exactly one)
expected_apn = "APN"
expected_amount = "Levy"
tolerance = 0.00                # per parcel; 0.00 means to the cent
fiscal_year = ""                # blank = the year before max_rates.fiscal_year

[backtest.inputs]               # file the workflow reads = last year's version of it
# "inputs/county_export.csv" = ".agent-grid/backtest/county_export.csv"
# "inputs/rates.csv" = ".agent-grid/backtest/rates.csv"

[backtest.known_differences]    # parcels the code shouldn't reproduce: last year's roll got
                                # them wrong, or someone decided them by hand. One line each,
                                # with the reason. Their amounts aren't compared; never list
                                # a whole district.
# "123-456-789" = "billed above the RMA max in 2025-26 (412.35 vs 412.3456)"
"""


class Unusable(Exception):
    """checks.toml or a file it names can't be used: the user's problem, not the analyst's."""


def load_config():
    if tomllib is None:
        raise Unusable("levy-check needs Python 3.11 or newer (for tomllib).")
    if not CONFIG.is_file():
        raise Unusable(f"No {CONFIG} here. Run: levy-check init")
    try:
        return tomllib.loads(CONFIG.read_text())
    except tomllib.TOMLDecodeError as e:
        raise Unusable(f"{CONFIG} isn't valid TOML: {e}")


def norm_apn(v):
    return re.sub(r"[\s.\-]", "", str(v or "")).upper()


def money(v):
    """'$1,203.44', '(12.00)', 1203.44 -> Decimal. Blank -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(repr(v))
    s = str(v).strip().replace("$", "").replace(",", "")
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(v)
    return -d if neg else d


def read_table(path, sheet=""):
    """Rows as dicts keyed by header. CSV or XLSX."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise Unusable("Reading .xlsx needs openpyxl (pip install openpyxl).")
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        out = [dict(zip(header, r)) for r in rows if any(c is not None and c != "" for c in r)]
        wb.close()
        return header, out
    with open(p, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        out = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
        return [h.strip() for h in (reader.fieldnames or [])], [{k.strip(): v for k, v in r.items() if k} for r in out]


def uncalculated(path, sheet, column):
    """True if an .xlsx column holds formulas that were never calculated."""
    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        i = header.index(column)
        found = any(isinstance(r[i], str) and r[i].startswith("=") for r in rows if len(r) > i)
        wb.close()
        return found
    except Exception:
        return False


def need_columns(header, cols, what):
    missing = [c for c in cols if c and c not in header]
    if missing:
        raise Unusable(f"{what} has no column(s) {', '.join(repr(c) for c in missing)}. "
                       f"Its columns are: {', '.join(header) or '(none)'}.")


def money_or_none(v):
    """Like money(), but an unreadable value counts as missing (it's reported elsewhere)."""
    try:
        return money(v)
    except ValueError:
        return None


def fmt(d):
    return "blank" if d is None else f"{d:,.2f}"


class Report:
    def __init__(self):
        self.lines, self.failed, self.failed_names = [], False, []

    def check(self, name, problems, detail=""):
        if problems:
            self.failed = True
            self.failed_names.append(name)
            shown = "; ".join(problems[:EXAMPLES]) + (f"; and {len(problems) - EXAMPLES} more" if len(problems) > EXAMPLES else "")
            self.lines.append(f"FAIL {name}: {detail or str(len(problems)) + ' found'}: {shown}")
        else:
            self.lines.append(f"pass {name}")

    def note(self, text):
        self.lines.append(text)

    def done(self):
        print("\n".join(self.lines))
        return 1 if self.failed else 0


def load_roll(cfg_roll, path=None):
    """Read the roll; returns list of (raw_apn, apn, amount, row) and problems parsing amounts."""
    path = path or cfg_roll.get("file", "")
    if not path:
        raise Unusable("checks.toml [roll] has no file.")
    apn_c, amt_c = cfg_roll.get("apn", "APN"), cfg_roll.get("amount", "Levy")
    try:
        header, rows = read_table(path, cfg_roll.get("sheet", ""))
    except FileNotFoundError:
        return None, [f"the workflow didn't write {path}"]
    need_columns(header, [apn_c, amt_c], path)
    parsed, bad = [], []
    for i, r in enumerate(rows, start=2):
        try:
            amt = money(r.get(amt_c))
        except ValueError:
            bad.append(f"row {i} APN {r.get(apn_c)}: {r.get(amt_c)!r} isn't an amount")
            continue
        parsed.append((r.get(apn_c), norm_apn(r.get(apn_c)), amt, r))
    return parsed, bad


# ---- Fiscal years and max rates ------------------------------------
def fiscal_start(v, what="a fiscal year"):
    """'2026-27', 'FY 2026-27', 'FY 26-27', '26/27', 2026 -> 2026."""
    if isinstance(v, int) and 1900 < v < 2200:
        return v
    t = str(v or "").strip()
    if not t:
        raise Unusable(f'{what} is blank: set it to a fiscal year like "2026-27".')
    m = re.search(r"(19|20)\d{2}", t)
    if m:
        return int(m.group(0))
    m = re.search(r"(?<!\d)(\d{2})\s*[-/]\s*\d{2}(?!\d)", t)
    if m:
        return 2000 + int(m.group(1))
    raise Unusable(f"{what} {v!r} isn't a fiscal year like \"2026-27\".")


def fy_label(y):
    return f"{y}-{(y + 1) % 100:02d}"


ROUNDING = {"half_up": ROUND_HALF_UP, "half_even": ROUND_HALF_EVEN, "down": ROUND_DOWN}


def max_rate_for(cls, spec, year, this_year):
    """The class's max rate per unit in fiscal year `year`, from checks.toml (never the code).
    None if checks.toml doesn't give one for that year."""
    for k, v in (spec.get("rates") or {}).items():
        if fiscal_start(k, f"max_rates.classes.{cls}.rates key") == year:
            return Decimal(str(v))
    if "rate" in spec and year == this_year:
        return Decimal(str(spec["rate"]))
    if "base_rate" in spec:
        if "base_year" not in spec:
            raise Unusable(f"max_rates.classes.{cls!r} has a base_rate but no base_year.")
        base, by = Decimal(str(spec["base_rate"])), fiscal_start(spec["base_year"], f"max_rates.classes.{cls}.base_year")
        esc = Decimal(str(spec.get("escalator", 0)))
        n = year - by
        if n < 0:
            return None
        mode = ROUNDING.get(spec.get("rounding", "half_up"))
        if mode is None:
            raise Unusable(f'max_rates.classes.{cls!r} rounding must be "half_up", "half_even", or "down".')
        places = spec.get("places", "")
        q = Decimal(1).scaleb(-int(places)) if str(places).strip() != "" else None
        r = base
        if spec.get("round_each_year") and q is not None:
            for _ in range(n):
                r = (r * (1 + esc)).quantize(q, rounding=mode)
        else:
            r = base * (1 + esc) ** n
            if q is not None:
                r = r.quantize(q, rounding=mode)
        return r
    return None


def show_rate(d):
    return f"{d.normalize():f}" if d == d.to_integral() else f"{d:f}".rstrip("0")


def check_max_rates(rep, cfg, rows, year, name, require_year=True, header=None):
    """Every parcel's levy is at most units x its class's max rate for `year`."""
    mr, roll = cfg.get("max_rates") or {}, cfg.get("roll") or {}
    classes = mr.get("classes") or {}
    this_year = fiscal_start(mr.get("fiscal_year"), "max_rates.fiscal_year")
    cc = mr.get("class_column", "")
    uc = mr.get("units_column", "") or roll.get("units", "")
    rc = roll.get("rate", "")
    if not cc and len(classes) != 1:
        raise Unusable("max_rates has more than one class, so it needs class_column: the roll column naming each parcel's class.")
    if header is not None:
        need_columns(header, [cc, uc], roll.get("file"))
    only = next(iter(classes)) if not cc else None
    over, rate_over, unknown, unreadable, skipped = [], [], [], [], set()
    for raw, _, a, r in rows:
        if a is None or a <= 0:
            continue
        cls = only if only is not None else str(r.get(cc) if r.get(cc) is not None else "").strip()
        spec = classes.get(cls)
        if spec is None:
            unknown.append(f"{raw} (class {cls or 'blank'!r})")
            continue
        m = max_rate_for(cls, spec, year, this_year)
        if m is None:
            skipped.add(cls)
            continue
        if uc:
            u = money_or_none(r.get(uc))
            if u is None:
                unreadable.append(f"{raw} units {r.get(uc)!r}")
                continue
        else:
            u = Decimal(1)
        limit = u * m
        if a > limit:
            over.append(f"{raw} ({cls}) levy {fmt(a)} > {show_rate(u)} x {show_rate(m)} = {show_rate(limit)}")
        if rc:
            rt = money_or_none(r.get(rc))
            if rt is not None and rt > m:
                rate_over.append(f"{raw} ({cls}) rate {show_rate(rt)} > max {show_rate(m)}")
    if skipped:
        if require_year:
            raise Unusable(f"max_rates gives no {fy_label(year)} max rate for class(es) {', '.join(sorted(skipped))}.")
        rep.note(f"  ({name}: no {fy_label(year)} max rate in checks.toml for {', '.join(sorted(skipped))}; those parcels weren't checked)")
    rep.check(f"{name}: every parcel's class has a max rate", unknown, f"{len(unknown)} parcel(s) with a class max_rates doesn't list")
    if uc:
        rep.check(f"{name}: units readable", unreadable)
    rep.check(f"{name}: levy within units x max rate", over,
              f"{len(over)} parcel(s) over the {fy_label(year)} maximum")
    if rc:
        rep.check(f"{name}: rate within max rate", rate_over)


# ---- This year's rules ---------------------------------------------
def run_rules(cfg, parsed, bad, header=None):
    roll, rules = cfg.get("roll", {}), cfg.get("rules", {})
    rep = Report()
    total = sum((a for *_, a, _ in parsed if a is not None), Decimal(0))
    rep.note(f"levy-check: {roll.get('file')}, {len(parsed)} parcels, total {fmt(total)}")
    rep.check("amounts readable", bad)
    blanks = [f"APN {raw}" for raw, _, a, _ in parsed if a is None]
    rep.check("no blank amounts", blanks)
    if header is None:
        header = list(parsed[0][3].keys()) if parsed else []

    if rules.get("no_blank_apn", True):
        rep.check("no_blank_apn", [f"row with amount {fmt(a)}" for _, apn, a, _ in parsed if not apn])
    unique_by = [c for c in (rules.get("unique_by") or []) if c]
    if unique_by:
        need_columns(header, unique_by, roll.get("file"))
        apn_c = roll.get("apn", "APN")
        seen = {}
        for raw, apn, _, r in parsed:
            k = tuple(apn if c == apn_c else str(r.get(c) if r.get(c) is not None else "").strip() for c in unique_by)
            if any(k):
                seen.setdefault(k, []).append(" / ".join(str(raw) if c == apn_c else x for c, x in zip(unique_by, k)))
        dups = [f"{v[0]} (x{len(v)})" for v in seen.values() if len(v) > 1]
        rep.check("unique_by", dups, f"{len(dups)} line(s) appear more than once by {' + '.join(unique_by)}")
    elif rules.get("unique_apn", True):
        seen = {}
        for raw, apn, _, _ in parsed:
            if apn:
                seen.setdefault(apn, []).append(raw)
        dups = [f"{v[0]} (x{len(v)})" for v in seen.values() if len(v) > 1]
        rep.check("unique_apn", dups, f"{len(dups)} APN(s) appear more than once")
    if rules.get("no_negative", True):
        rep.check("no_negative", [f"{raw} {fmt(a)}" for raw, _, a, _ in parsed if a is not None and a < 0])
    if rules.get("whole_cents", True):
        rep.check("whole_cents", [f"{raw} {a}" for raw, _, a, _ in parsed if a is not None and a != a.quantize(Decimal("0.01"))])

    mt_c, u_c, r_c = roll.get("max_tax", ""), roll.get("units", ""), roll.get("rate", "")
    if rules.get("max_tax") or rules.get("max_tax_is_units_x_rate"):
        need_columns(header, [mt_c] + ([u_c, r_c] if rules.get("max_tax_is_units_x_rate") else []), roll.get("file"))
    if rules.get("max_tax"):
        if not mt_c:
            raise Unusable("rules.max_tax is on, but [roll] max_tax names no column.")
        over = []
        for raw, _, a, r in parsed:
            m = money_or_none(r.get(mt_c))
            if a is not None and m is not None and a > m:
                over.append(f"{raw} levy {fmt(a)} > max {fmt(m)}")
        rep.check("max_tax", over, f"{len(over)} parcel(s) over the maximum")
    if rules.get("max_tax_is_units_x_rate"):
        mode = {"half_up": ROUND_HALF_UP, "half_even": ROUND_HALF_EVEN}.get(rules.get("rounding", "half_up"))
        if mode is None:
            raise Unusable('rules.rounding must be "half_up" or "half_even".')
        wrong = []
        for raw, _, _, r in parsed:
            m, u, rt = money_or_none(r.get(mt_c)), money_or_none(r.get(u_c)), money_or_none(r.get(r_c))
            if None in (m, u, rt):
                continue
            want = (u * rt).quantize(Decimal("0.01"), rounding=mode)
            if m != want:
                wrong.append(f"{raw} max {fmt(m)} but units x rate = {fmt(want)}")
        rep.check("max_tax_is_units_x_rate", wrong)

    mr = cfg.get("max_rates") or {}
    if mr.get("classes"):
        year = fiscal_start(mr.get("fiscal_year"), "max_rates.fiscal_year")
        check_max_rates(rep, cfg, parsed, year, "max_rates", require_year=True, header=header)

    want_total = rules.get("total", 0)
    if want_total:
        want, tol = Decimal(str(want_total)), Decimal(str(rules.get("total_tolerance", 1.00)))
        diff = total - want
        rep.check("total", [f"expected {fmt(want)}, got {fmt(total)} (off by {fmt(diff)})"] if abs(diff) > tol else [])

    tot = cfg.get("totals") or {}
    gc = tot.get("column", "")
    if gc and tot.get("expected"):
        need_columns(header, [gc], roll.get("file"))
        tol = Decimal(str(tot.get("tolerance", 1.00)))
        sums = {}
        for _, _, a, r in parsed:
            g = str(r.get(gc) if r.get(gc) is not None else "").strip()
            sums[g] = sums.get(g, Decimal(0)) + (a or 0)
        off = []
        for g, w in tot["expected"].items():
            want, have = Decimal(str(w)), sums.get(str(g).strip(), Decimal(0))
            if abs(have - want) > tol:
                off.append(f"{g}: expected {fmt(want)}, got {fmt(have)} (off by {fmt(have - want)})")
        rep.check(f"totals by {gc}", off, f"{len(off)} group(s) off")
        extra = sorted(g for g, v in sums.items() if v and g not in {str(k).strip() for k in tot["expected"]})
        if extra:
            rep.note(f"  (groups with levy but no expected total in [totals.expected]: {', '.join(repr(g) for g in extra[:EXAMPLES])})")

    pf = rules.get("parcels_file", "")
    if pf:
        want, _ = parcel_list(rules)
        have = {apn: raw for raw, apn, _, _ in parsed if apn}
        rep.check("every input parcel in the roll", [str(want[a]) for a in want if a not in have],
                  f"{sum(1 for a in want if a not in have)} parcel(s) in {pf} missing from the roll")
        rep.check("no parcel outside the input", [str(have[a]) for a in have if a not in want],
                  f"{sum(1 for a in have if a not in want)} parcel(s) in the roll aren't in {pf}")
    return rep


_PARCELS = {}


def parcel_list(rules):
    pf = rules.get("parcels_file", "")
    if pf not in _PARCELS:
        try:
            ph, prows = read_table(pf)
        except FileNotFoundError:
            raise Unusable(f"rules.parcels_file {pf} doesn't exist.")
        pa = rules.get("parcels_apn", "APN")
        need_columns(ph, [pa], pf)
        _PARCELS[pf] = ({norm_apn(r.get(pa)): r.get(pa) for r in prows if norm_apn(r.get(pa))}, ph)
    return _PARCELS[pf]


def cmd_rules():
    cfg = load_config()
    roll = cfg.get("roll", {})
    parsed, bad = load_roll(roll)
    if parsed is None:
        rep = Report()
        rep.check("roll written", bad, "missing")
        return rep.done()
    header, _ = read_table(roll.get("file"), roll.get("sheet", "")) if not parsed else (list(parsed[0][3].keys()), None)
    rep = run_rules(cfg, parsed, bad, header)
    if any(a is None for _, _, a, _ in parsed) and uncalculated(roll.get("file", ""), roll.get("sheet", ""), roll.get("amount", "Levy")):
        rep.note("  (the amount column holds formulas with no saved results: the workbook was written by a "
                 "script and never calculated. Write values, or open and save it in Excel, before checking)")
    return rep.done()


# ---- Backtest --------------------------------------------------------
def cmd_backtest_inputs():
    bt = load_config().get("backtest", {})
    exp = (bt.get("expected") or "").strip()
    if not exp:
        return 0                                    # no backtest set up
    if not Path(exp).is_file():
        raise Unusable(f"backtest.expected {exp!r} doesn't exist. Put last year's certified roll there.")
    wf = (bt.get("workflow") or "").strip()
    if not wf:
        found = sorted(p.parent.name for p in Path(".agent-grid/workflows").glob("*/run.sh"))
        if not found:
            return 0                                # starts once the analyst writes a workflow
        if len(found) > 1:
            raise Unusable(f"backtest.workflow is blank, but there are {len(found)} workflows "
                           f"({', '.join(found)}). Name the one that writes the roll.")
        wf = found[0]
    elif not Path(f".agent-grid/workflows/{wf}/run.sh").is_file():
        raise Unusable(f"backtest.workflow {wf!r} isn't a workflow in .agent-grid/workflows/.")
    print(wf)
    for dest, src in (bt.get("inputs") or {}).items():
        if not Path(src).is_file():
            raise Unusable(f"backtest input {src!r} (for {dest}) doesn't exist.")
        print(f"{dest}\t{src}")
    return 0


def load_expected(bt):
    exp_cfg = {"file": bt.get("expected"), "apn": bt.get("expected_apn", "APN"), "amount": bt.get("expected_amount", "Levy")}
    exp, ebad = load_roll(exp_cfg)
    if exp is None or ebad:
        raise Unusable(f"Can't read the certified roll {bt.get('expected')}: {'; '.join((ebad or [])[:3])}")
    return exp


def backtest_year(cfg):
    bt, mr = cfg.get("backtest", {}), cfg.get("max_rates") or {}
    if str(bt.get("fiscal_year") or "").strip():
        return fiscal_start(bt["fiscal_year"], "backtest.fiscal_year")
    return fiscal_start(mr.get("fiscal_year"), "max_rates.fiscal_year") - 1


def compare(cfg, got, bad, exp, check_max=True):
    bt = cfg.get("backtest", {})
    rep = Report()
    tol = Decimal(str(bt.get("tolerance", 0)))
    known = {norm_apn(k): v for k, v in (bt.get("known_differences") or {}).items()}

    def by_apn(rows):
        d, dup = {}, []
        for raw, apn, a, r in rows:
            if apn in d:
                dup.append(str(raw))
                d[apn] = (raw, (d[apn][1] or 0) + (a or 0), d[apn][2])
            else:
                d[apn] = (raw, a, r)
        return d, dup

    g, gdup = by_apn(got)
    e, _ = by_apn(exp)
    gt = sum((a for _, a, _ in g.values() if a is not None), Decimal(0))
    et = sum((a for _, a, _ in e.values() if a is not None), Decimal(0))
    rep.note(f"levy-check backtest: last year's inputs through this code, vs {bt.get('expected')}")
    rep.note(f"  {len(g)} parcels, total {fmt(gt)}; certified: {len(e)} parcels, total {fmt(et)}; difference {fmt(gt - et)}")
    rep.check("amounts readable", bad)
    unique_by = [c for c in ((cfg.get("rules") or {}).get("unique_by") or []) if c]
    if not unique_by:                               # several lines per APN are summed when unique_by is set
        rep.check("no duplicate APNs", gdup)
    rep.check("same parcels as certified", [str(e[a][0]) for a in e if a not in g],
              f"{sum(1 for a in e if a not in g)} certified parcel(s) missing")
    rep.check("no extra parcels", [str(g[a][0]) for a in g if a not in e],
              f"{sum(1 for a in g if a not in e)} parcel(s) not on the certified roll")
    diffs, excused = [], []
    for a, (raw, want, _) in e.items():
        if a in g:
            have = g[a][1]
            if (have is None) != (want is None) or (have is not None and abs(have - want) > tol):
                (excused if a in known else diffs).append(f"{raw} certified {fmt(want)}, got {fmt(have)}")
    rep.check("same amounts as certified", diffs, f"{len(diffs)} parcel(s) differ")
    if known:
        missing = [k for k in known if k not in e]
        rep.note(f"  ({len(known)} known difference(s) in checks.toml aren't compared"
                 f"{'; ' + str(len(excused)) + ' of them differ this run' if excused else ''}"
                 f"{'; not on the certified roll: ' + ', '.join(missing[:EXAMPLES]) if missing else ''})")

    mr = cfg.get("max_rates") or {}
    if check_max and mr.get("classes"):
        year = backtest_year(cfg)
        check_max_rates(rep, cfg, got, year, "backtest max_rates", require_year=False)
        # Last year's certified roll over its own max: those parcels can't be required to match
        joined = [(e[a][0], a, e[a][1], g[a][2]) for a in e if a in g and a not in known]
        probe = Report()
        check_max_rates(probe, cfg, joined, year, "certified", require_year=False)
        over = [ln for ln in probe.lines if ln.startswith("FAIL certified: levy within")]
        if over:
            rep.note("  NOTE last year's certified roll is over its own max for some parcels (amounts from the "
                     "certified roll, classes and units from this run): " + over[0].split(": ", 2)[-1]
                     + ". List them under [backtest.known_differences] so the code isn't required to repeat it.")
    return rep


def cmd_backtest_compare():
    cfg = load_config()
    roll, bt = cfg.get("roll", {}), cfg.get("backtest", {})
    got, bad = load_roll(roll)
    if got is None:
        rep = Report()
        rep.check("backtest roll written", bad, "missing")
        return rep.done()
    return compare(cfg, got, bad, load_expected(bt)).done()


# ---- Hand-calculated parcels (kept outside the folder) ---------------
def cmd_spot(path):
    path = path or os.environ.get("AGENT_GRID_SPOT", "")
    if not path:
        raise Unusable("No spot-check file named (levy-check spot FILE, or AGENT_GRID_SPOT).")
    if not Path(path).is_file():
        raise Unusable(f"The spot-check file {path} doesn't exist.")
    try:
        spot = tomllib.loads(Path(path).read_text())
    except tomllib.TOMLDecodeError as e:
        raise Unusable(f"The spot-check file isn't valid TOML: {e}")
    want = spot.get("parcels") or {}
    cfg = load_config()
    got, bad = load_roll(cfg.get("roll", {}))
    rep = Report()
    if got is None:
        rep.check("roll written", bad, "missing")
        return rep.done()
    if not want:
        rep.note("levy-check spot: the spot-check file lists no parcels yet")
        return rep.done()
    have = {}
    for _, apn, a, _ in got:                        # several lines for one APN are added up
        have[apn] = (have[apn] or 0) + (a or 0) if apn in have else a
    tol = Decimal(str(spot.get("tolerance", 0)))
    off = []
    for k, v in want.items():
        w, a = Decimal(str(v)), norm_apn(k)
        if a not in have:
            off.append(f"{k}: expected {fmt(w)}, not on the roll")
        elif have[a] is None or abs(have[a] - w) > tol:
            off.append(f"{k}: expected {fmt(w)}, got {fmt(have[a])}")
    rep.note(f"levy-check spot: {len(want)} hand-calculated parcel(s)"
             f"{' for ' + str(spot['fiscal_year']) if spot.get('fiscal_year') else ''}")
    rep.check("hand-calculated parcels match", off, f"{len(off)} of {len(want)} don't match")
    return rep.done()


# ---- Self-test: plant errors and see what catches them ---------------
def _with(row_t, amount=None, apn=None, cols=None, amt_c="Levy", apn_c="APN"):
    raw, n, a, r = row_t
    r = dict(r)
    if amount is not None:
        a = amount
        r[amt_c] = str(amount)
    if apn is not None:
        raw, n = apn, norm_apn(apn)
        r[apn_c] = apn
    for k, v in (cols or {}).items():
        r[k] = v
    return (raw, n, a, r)


def _cent(d):
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def roll_mutations(cfg, rows):
    """(label, mutated rows or None, why not applicable)"""
    roll = cfg.get("roll", {})
    amt_c, apn_c = roll.get("amount", "Levy"), roll.get("apn", "APN")
    pos = [i for i, (_, apn, a, _) in enumerate(rows) if apn and a is not None and a > 0]
    if not pos:
        return []
    i = pos[len(pos) // 2]
    t = rows[i]
    a = t[2]
    w = lambda **k: _with(t, amt_c=amt_c, apn_c=apn_c, **k)
    out = [
        ("one parcel left off the roll", rows[:i] + rows[i + 1:], ""),
        ("one parcel billed twice", rows + [t], ""),
        ("a parcel added that isn't in the district", rows + [w(apn="999-999-999")], ""),
        ("one parcel's levy 10% too high", rows[:i] + [w(amount=_cent(a * Decimal("1.10")))] + rows[i + 1:], ""),
        ("every levy 2% too high (an escalator applied twice)",
         [_with(x, amount=_cent(x[2] * Decimal("1.02")), amt_c=amt_c, apn_c=apn_c) if x[2] is not None else x for x in rows], ""),
        ("a fraction of a cent on one parcel", rows[:i] + [w(amount=a + Decimal("0.001"))] + rows[i + 1:], ""),
        ("a negative levy", rows[:i] + [w(amount=-a)] + rows[i + 1:], ""),
        ("a parcel with a blank APN", rows[:i] + [w(apn="")] + rows[i + 1:], ""),
    ]
    # Over the maximum, by whichever max check is set up
    limit = None
    mr = cfg.get("max_rates") or {}
    if mr.get("classes"):
        year = fiscal_start(mr.get("fiscal_year"), "max_rates.fiscal_year")
        cc, uc = mr.get("class_column", ""), mr.get("units_column", "") or roll.get("units", "")
        classes = mr["classes"]
        cls = next(iter(classes)) if not cc else str(t[3].get(cc) or "").strip()
        if cls in classes:
            m = max_rate_for(cls, classes[cls], year, year)
            u = money_or_none(t[3].get(uc)) if uc else Decimal(1)
            if m is not None and u is not None:
                limit = u * m
    if limit is None and (cfg.get("rules") or {}).get("max_tax") and roll.get("max_tax"):
        limit = money_or_none(t[3].get(roll["max_tax"]))
    if limit is not None:
        over = limit.quantize(Decimal("0.01"), rounding=ROUND_UP)
        over = over + Decimal("0.01") if over == limit else over
        out.append(("one parcel a cent over its maximum", rows[:i] + [w(amount=over)] + rows[i + 1:], ""))
    else:
        out.append(("one parcel a cent over its maximum", None, "no max check is set up ([max_rates] or rules.max_tax)"))
    # Wrong rate class: move the parcel to the class with the lowest max rate
    if mr.get("classes") and mr.get("class_column") and len(mr["classes"]) > 1:
        year = fiscal_start(mr.get("fiscal_year"), "max_rates.fiscal_year")
        cc = mr["class_column"]
        cur = str(t[3].get(cc) or "").strip()
        rated = [(max_rate_for(c, s, year, year), c) for c, s in mr["classes"].items() if c != cur]
        rated = [x for x in rated if x[0] is not None]
        if rated:
            low = min(rated)[1]
            out.append((f"one parcel in the wrong rate class ({cur or 'blank'} -> {low})",
                        rows[:i] + [w(cols={cc: low})] + rows[i + 1:], ""))
    return out


def backtest_mutations(cfg, rows):
    bt = cfg.get("backtest", {})
    known = {norm_apn(k) for k in (bt.get("known_differences") or {})}
    pos = [i for i, (_, apn, a, _) in enumerate(rows) if apn and apn not in known and a is not None and a > 0]
    out = []
    if pos:
        i = pos[len(pos) // 2]
        t = rows[i]
        out += [
            ("one parcel off by a cent", rows[:i] + [_with(t, amount=t[2] + Decimal("0.01"))] + rows[i + 1:], ""),
            ("one parcel left off", rows[:i] + rows[i + 1:], ""),
            ("one extra parcel", rows + [_with(t, apn="999-999-999")], ""),
        ]
    kpos = [i for i, (_, apn, a, _) in enumerate(rows) if apn in known and a is not None]
    if kpos:
        i = kpos[0]
        t = rows[i]
        out.append((f"a known-difference parcel ({t[0]}) off by $10", rows[:i] + [_with(t, amount=t[2] + 10)] + rows[i + 1:], ""))
    return out


def cmd_selftest():
    cfg = load_config()
    roll, bt = cfg.get("roll", {}), cfg.get("backtest", {})
    print("levy-check selftest: plants errors in copies of the rolls and reports which checks catch each.")
    print("Nothing on disk changes. A MISSED error is a gap: the manager and your own review have to cover it.")
    missed = total = 0
    parsed, bad = load_roll(roll)
    if parsed is None:
        print(f"\nThis year's roll: there's no {roll.get('file')} yet, so its rules weren't tested. "
              "Run selftest again once a roll exists.")
    else:
        header = list(parsed[0][3].keys()) if parsed else []
        base = run_rules(cfg, parsed, bad, header)
        before = set(base.failed_names)
        print(f"\nThis year's roll ({roll.get('file')}, {len(parsed)} lines):")
        if before:
            print(f"  (already failing on this roll, so not counted: {', '.join(sorted(before))})")
        for label, rows, why in roll_mutations(cfg, parsed):
            total += 1
            if rows is None:
                missed += 1
                print(f"  MISSED  {label}  ({why})")
                continue
            rep = run_rules(cfg, rows, bad, header)
            by = [n for n in rep.failed_names if n not in before]
            if by:
                print(f"  caught  {label}  (by {', '.join(by)})")
            else:
                missed += 1
                print(f"  MISSED  {label}")
    if (bt.get("expected") or "").strip():
        exp = load_expected(bt)
        base = compare(cfg, exp, [], exp, check_max=False)
        before = set(base.failed_names)
        print(f"\nBacktest (last year's certified roll {bt.get('expected')}, as if the code reproduced it exactly):")
        for label, rows, _ in backtest_mutations(cfg, exp):
            total += 1
            rep = compare(cfg, rows, [], exp, check_max=False)
            by = [n for n in rep.failed_names if n not in before]
            if by:
                print(f"  caught  {label}  (by {', '.join(by)})")
            else:
                missed += 1
                note = "  (known differences aren't compared)" if label.startswith("a known-difference") else ""
                print(f"  MISSED  {label}{note}")
    else:
        print("\nBacktest: not set up, so nothing compares results with last year's certified roll.")
    print(f"\n{missed} of {total} planted errors got through." if total else "\nNothing to test.")
    return 0


def main(argv):
    cmd = argv[1] if len(argv) > 1 else "rules"
    if any(a in ("-h", "--help", "help") for a in argv[1:]):
        print(__doc__.strip())
        return 0
    try:
        if cmd == "init":
            if CONFIG.exists():
                print(f"{CONFIG} already exists; not overwriting it.")
                return 2
            CONFIG.parent.mkdir(parents=True, exist_ok=True)
            CONFIG.write_text(TEMPLATE)
            print(f"Wrote {CONFIG}. Fill in the column names and turn on the rules that apply,")
            print("then put last year's certified roll and inputs in .agent-grid/backtest/.")
            return 0
        if cmd == "rules":
            return cmd_rules()
        if cmd == "selftest":
            return cmd_selftest()
        if cmd == "spot":
            return cmd_spot(argv[2] if len(argv) > 2 else "")
        if cmd == "roll-path":
            print(load_config().get("roll", {}).get("file", ""))
            return 0
        if cmd == "backtest-inputs":
            return cmd_backtest_inputs()
        if cmd == "backtest-compare":
            return cmd_backtest_compare()
        print(__doc__.strip())
        return 2
    except Unusable as e:
        print(f"checks.toml problem: {e}")
        return 2
    except ValueError as e:
        print(f"checks.toml problem: a number in it isn't a number ({e}).")
        return 2


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv))
    except BrokenPipeError:     # output piped into head and cut short
        sys.exit(0)
